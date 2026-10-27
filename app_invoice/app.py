"""Application web locale (Flask) : on dépose des factures PDF (glisser-déposer), chacune est lue par Docling puis par
le LLM local (Ollama, dans Docker), et le résultat s'affiche dans un tableau modifiable, à enregistrer ou à exporter
en Excel. Même prompt simple et même schéma que llm_local/test_llm_ollama.py.

Aucun masquage : rien ne quitte le serveur. Application autonome : n'importe aucun autre fichier du projet
(seulement ses propres templates/index.html, static/style.css et static/app.js).
Les factures sont traitées une par une (le LLM occupe tout le processeur) : les suivantes attendent leur tour.
Fichiers : Output/app/<date>_<facture>.pdf (le PDF déposé) et .json (le résultat, avec les corrections enregistrées).

Usage (depuis le dossier du projet, avec l'environnement Docling ; Ollama doit tourner) :
    .venv-docling/bin/python app_invoice/app.py                    -> http://localhost:8000
    .venv-docling/bin/python app_invoice/app.py -m gemma3:4b --port 8000
"""
import argparse
import io
import json
import logging
import os
import queue
import re
import threading
import time
import urllib.error
import urllib.request
import uuid
from datetime import datetime
from pathlib import Path

from docling.datamodel.base_models import InputFormat
from docling.datamodel.pipeline_options import PdfPipelineOptions
from docling.document_converter import DocumentConverter, PdfFormatOption
from docling_core.types.doc import ContentLayer
from flask import Flask, abort, jsonify, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

PROJET = Path(__file__).resolve().parents[1]  # dossier facture/ : Output
SORTIE = PROJET / "Output" / "app"
logging.getLogger("docling").setLevel(logging.WARNING)

# ---------------------------------------------------------------------------------------------------------------
# Champs à extraire (identiques à test_llm_ollama.py) : nom -> (type, explication donnée au modèle)
# Le tableau de la page est construit à partir d'eux : rien à changer dans la page pour d'autres champs.
# ---------------------------------------------------------------------------------------------------------------
TYPE_DOCUMENT = "une facture"
CHAMPS = {
    "fournisseur": ("texte", "société qui émet la facture et reçoit le paiement (NETIS est le client, jamais le "
                             "fournisseur), sans forme juridique (Inc, Ltd, LLC, SAS...)"),
    "date_paiement": ("texte", "date de paiement, sinon date d'échéance, sinon date de la facture jamais une date de début ou de fin de période d’abonnement , au format JJ/MM/AAAA"),
    "devise": ("texte", "code ISO à 3 lettres (USD, EUR, MAD...)"),
    "total": ("nombre", "montant total écrit sur le document, taxes comprises"),
}
LISTE = ("articles", "une entrée par ligne d'article des tableaux, dans l'ordre ; pas les sous-totaux, taxes, totaux "
                     "ni paiements", {
    "description": ("texte", "nom de l'article tel qu'écrit"),
    "quantite": ("nombre?", "quantité écrite, sinon null"),
    "prix_unitaire": ("nombre?", "prix unitaire écrit, sinon null"),
    "prix": ("nombre", "montant de la ligne (négatif s'il est écrit avec un signe -)"),
})
TYPES_JSON = {"texte": {"type": "string"}, "nombre": {"type": "number"},
              "texte?": {"type": ["string", "null"]}, "nombre?": {"type": ["number", "null"]}}


def construire_schema():
    proprietes = {nom: TYPES_JSON[t] for nom, (t, _) in CHAMPS.items()}
    nom, _, champs = LISTE
    proprietes[nom] = {"type": "array", "items": {"type": "object",
                                                   "properties": {n: TYPES_JSON[t] for n, (t, _) in champs.items()},
                                                   "required": list(champs)}}
    return {"type": "object", "properties": proprietes, "required": list(proprietes)}


def construire_prompt():
    lignes = [f"Tu extrais les données d'un document : {TYPE_DOCUMENT}. Il est donné en Markdown : les tableaux sont "
              "écrits avec des colonnes | ... |.",
              "Remplis le JSON uniquement avec ce qui est écrit dans le document : recopie les textes et les nombres, "
              "ne calcule rien, n'invente rien. Une information absente vaut null.",
              "Montants : nombres à point décimal, sans devise ni séparateur de milliers (\"1 234,56 €\" = 1234.56).",
              ""]
    lignes += [f"- {nom} : {explication}" for nom, (_, explication) in CHAMPS.items()]
    nom, explication, champs = LISTE
    lignes.append(f"- {nom} : {explication}")
    lignes += [f"  - {n} : {e}" for n, (_, e) in champs.items()]
    return "\n".join(lignes)


# ---------------------------------------------------------------------------------------------------------------
# Lecture du PDF (Docling) et appel du LLM local (Ollama)
# ---------------------------------------------------------------------------------------------------------------
def convertisseur():
    options = PdfPipelineOptions()
    options.do_ocr = False                   # PDF numériques : le texte est déjà dans le fichier
    options.do_table_structure = True        # vrais tableaux (cellules sur plusieurs lignes, colonnes nommées)
    options.generate_picture_images = False  # pas de logos en base64 dans le texte
    return DocumentConverter(format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=options)})


def pdf_en_markdown(conv, pdf):
    # tableaux compacts : moins de tokens à lire ; en-têtes et pieds de page gardés : le nom du fournisseur n'est
    # parfois écrit que là (OVH, mentions légales)
    return conv.convert(str(pdf)).document.export_to_markdown(
        compact_tables=True, escape_html=False, escape_underscores=False,
        included_content_layers={ContentLayer.BODY, ContentLayer.FURNITURE})


def appeler_ollama(url, modele, prompt, schema, texte):
    corps = {
        "model": modele,
        "messages": [{"role": "system", "content": prompt}, {"role": "user", "content": texte}],
        "format": schema,          # réponse forcée au format du schéma
        "stream": False,
        "think": False,            # pas de raisonnement écrit avant la réponse (trop lent sur CPU)
        "keep_alive": "30m",       # le modèle reste chargé entre deux factures
        "options": {"temperature": 0, "num_ctx": 8192},
    }
    req = urllib.request.Request(f"{url}/api/chat", data=json.dumps(corps).encode(),
                                 headers={"Content-Type": "application/json"})
    try:
        with urllib.request.urlopen(req, timeout=3600) as r:
            rep = json.load(r)
    except urllib.error.HTTPError as e:
        raise RuntimeError(f"Erreur d'Ollama ({e.code}) : {e.read().decode(errors='replace')[:500]}")
    except urllib.error.URLError as e:
        raise RuntimeError(f"Ollama injoignable sur {url} ({e.reason}) : le conteneur ollama tourne-t-il ?")
    return json.loads(rep["message"]["content"])


# ---------------------------------------------------------------------------------------------------------------
# File d'attente : les factures sont traitées une par une, en arrière-plan
# ---------------------------------------------------------------------------------------------------------------
app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 50 * 1024 * 1024  # 50 Mo par envoi
factures = {}                # id -> facture déposée (statut, résultat...), dans l'ordre de dépôt
file_attente = queue.Queue()
reglages = {"modele": "gemma3:4b", "url": "http://localhost:11434"}


def travailleur():
    conv, prompt, schema = convertisseur(), construire_prompt(), construire_schema()
    while True:
        f = file_attente.get()
        f["debut"] = time.time()
        try:
            f["statut"] = "lecture"
            texte = pdf_en_markdown(conv, f["pdf"])                                       # 1. Docling
            f["statut"] = "llm"
            f["resultat"] = appeler_ollama(reglages["url"], reglages["modele"], prompt, schema, texte)  # 2. LLM
            f["statut"] = "fini"
            enregistrer(f)
        except Exception as e:
            f["statut"], f["erreur"] = "erreur", str(e) or e.__class__.__name__
        f["fin"] = time.time()


def enregistrer(f):
    f["pdf"].with_suffix(".json").write_text(json.dumps(f["resultat"], indent=2, ensure_ascii=False), encoding="utf-8")


def public(f):
    """Ce que la page reçoit pour une facture."""
    attente = [g["id"] for g in factures.values() if g["statut"] == "attente"]
    return {"id": f["id"], "nom": f["nom"], "statut": f["statut"], "resultat": f["resultat"], "erreur": f["erreur"],
            "modifie": f["modifie"], "rang": attente.index(f["id"]) + 1 if f["id"] in attente else None,
            "duree": round((f["fin"] or time.time()) - f["debut"]) if f["debut"] else None}


def classeur_excel(liste):
    """Une feuille des factures et une feuille des lignes (articles) ; en rouge si la somme des lignes ≠ total."""
    nom_liste, _, champs_liste = LISTE
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Factures"
    feuille.append(["Fichier", *CHAMPS, f"Somme des {nom_liste}", "Contrôle"])
    lignes = classeur.create_sheet(nom_liste.capitalize())
    lignes.append(["Fichier", *champs_liste])
    rouge = PatternFill("solid", fgColor="F8D7DA")
    for f in liste:
        donnees = f.get("donnees") or {}
        articles = donnees.get(nom_liste) or []
        somme = round(sum(a.get("prix") or 0 for a in articles), 2)
        ok = donnees.get("total") is None or abs(somme - donnees["total"]) < 0.01
        feuille.append([f["nom"], *[donnees.get(c) for c in CHAMPS], somme, "OK" if ok else "À vérifier"])
        if not ok:
            for cellule in feuille[feuille.max_row]:
                cellule.fill = rouge
        for a in articles:
            lignes.append([f["nom"], *[a.get(c) for c in champs_liste]])
    for f in (feuille, lignes):
        for cellule in f[1]:
            cellule.font = Font(bold=True)
        for colonne in f.columns:
            largeur = max(len(str(c.value or "")) for c in colonne)
            f.column_dimensions[colonne[0].column_letter].width = min(60, max(10, largeur + 2))
    tampon = io.BytesIO()
    classeur.save(tampon)
    tampon.seek(0)
    return tampon


# ---------------------------------------------------------------------------------------------------------------
# Pages et API
# ---------------------------------------------------------------------------------------------------------------
@app.get("/")
def accueil():
    return render_template("index.html", modele=reglages["modele"])


@app.get("/api/champs")
def champs():
    """Les champs à extraire : la page construit le tableau à partir d'eux."""
    nom, aide, champs_liste = LISTE
    return jsonify(modele=reglages["modele"],
                   champs=[{"nom": n, "type": t, "aide": a} for n, (t, a) in CHAMPS.items()],
                   liste={"nom": nom, "aide": aide,
                          "champs": [{"nom": n, "type": t, "aide": a} for n, (t, a) in champs_liste.items()]})


@app.get("/api/factures")
def liste_factures():
    return jsonify([public(f) for f in factures.values()])


@app.post("/api/factures")
def deposer():
    """Un ou plusieurs PDF (champ "pdf" du formulaire) mis dans la file d'attente."""
    deposees = []
    for fichier in request.files.getlist("pdf"):
        contenu = fichier.read()
        if not contenu.startswith(b"%PDF"):
            return jsonify(erreur=f"{fichier.filename} n'est pas un PDF"), 400
        nom = Path(fichier.filename or "facture.pdf").name
        pdf = SORTIE / f"{datetime.now():%Y%m%d-%H%M%S}_{re.sub(r'[^\w.-]+', '_', Path(nom).stem)[:80]}.pdf"
        pdf.write_bytes(contenu)
        f = {"id": uuid.uuid4().hex[:12], "nom": nom, "pdf": pdf, "statut": "attente", "resultat": None,
             "erreur": None, "modifie": False, "debut": None, "fin": None}
        factures[f["id"]] = f
        file_attente.put(f)
        deposees.append(public(f))
    return jsonify(deposees)


@app.post("/api/factures/<id_>")
def corriger(id_):
    """Corrections faites dans le tableau : elles remplacent le résultat du LLM (et son fichier .json)."""
    if id_ not in factures:
        abort(404)
    f = factures[id_]
    f["resultat"], f["modifie"] = request.get_json(), True
    enregistrer(f)
    return jsonify(public(f))


@app.get("/api/factures/<id_>/pdf")
def voir_pdf(id_):
    if id_ not in factures:
        abort(404)
    return send_file(factures[id_]["pdf"], mimetype="application/pdf", download_name=factures[id_]["nom"])


@app.post("/api/excel")
def exporter():
    """[{nom, donnees}] (les tableaux tels qu'affichés) -> classeur Excel."""
    return send_file(classeur_excel(request.get_json()), as_attachment=True,
                     download_name=f"factures_{datetime.now():%Y%m%d-%H%M}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("-m", "--modele", default=os.environ.get("MODELE", reglages["modele"]))
    ap.add_argument("--url", default=os.environ.get("OLLAMA_URL", reglages["url"]))
    ap.add_argument("--port", type=int, default=8000)
    ap.add_argument("--host", default="127.0.0.1", help="127.0.0.1 : accessible depuis ce serveur seulement")
    args = ap.parse_args()
    reglages.update(modele=args.modele, url=args.url)
    SORTIE.mkdir(parents=True, exist_ok=True)

    # le modèle est-il installé dans Ollama ?
    try:
        with urllib.request.urlopen(f"{args.url}/api/tags", timeout=5) as r:
            modeles = [m["name"] for m in json.load(r).get("models", [])]
        if args.modele not in modeles and f"{args.modele}:latest" not in modeles:
            print(f"⚠ Modèle {args.modele} absent d'Ollama (installés : {', '.join(modeles) or 'aucun'}) :\n"
                  f"   docker exec ollama ollama pull {args.modele}")
    except OSError:
        print(f"⚠ Ollama injoignable sur {args.url} : le conteneur ollama tourne-t-il ?")

    threading.Thread(target=travailleur, daemon=True).start()
    print(f"Application : http://localhost:{args.port}   (modèle {args.modele}, Ctrl+C pour arrêter)")
    app.run(host=args.host, port=args.port, threaded=True)


if __name__ == "__main__":
    main()
