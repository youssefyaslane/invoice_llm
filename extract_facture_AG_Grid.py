"""Extraction des informations d'une facture AG Grid (PDF) vers Excel.

Usage :
    python extract_facture.py [facture.pdf ...] [-o dossier_sortie]

Sans argument, traite Data/INV-55868.pdf.
"""
import argparse
import re
from datetime import datetime
from pathlib import Path

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent
NUM = r"-?[\d,]+\.\d{2}"

# Ligne article : code, début de description, quantité, prix unitaire, TVA, montant
LIGNE_RE = re.compile(
    rf"^(?P<item>[A-Z0-9]{{6,}})\s+(?P<desc>.+?)\s+(?P<qte>{NUM})\s+"
    rf"(?P<pu>{NUM})\s+(?P<tva>.+?)\s+(?P<montant>{NUM})$"
)


def to_float(s):
    return float(s.replace(",", "")) if s else None


def to_date(s):
    try:
        return datetime.strptime(s.strip(), "%d %b %Y").date()
    except (ValueError, AttributeError):
        return s


def search(pattern, text, flags=0):
    m = re.search(pattern, text, flags)
    return m.group(1).strip() if m else None


def zone_lignes(page, x0, x1, top, bottom):
    """Texte d'une zone rectangulaire de la page, ligne par ligne."""
    zone = page.crop((x0, top, x1, bottom)).extract_text(x_tolerance=1) or ""
    return [l.strip() for l in zone.splitlines() if l.strip()]


def extraire_entete(page):
    w = page.width
    # Repère vertical : la ligne d'en-tête du tableau ("Item Description ...")
    titres = page.search("Item") or []
    bas = titres[0]["top"] - 5 if titres else 300

    # Bloc client (gauche, sous "TAX INVOICE")
    titre = page.search("INVOICE")
    haut_client = titre[0]["bottom"] + 2 if titre else 140
    client = zone_lignes(page, 0, w * 0.57, haut_client, bas)

    # Bloc méta (Invoice Date / Number / Reference / VAT Number)
    meta_lignes = zone_lignes(page, w * 0.57, w * 0.76, 0, bas)
    meta = {}
    labels = ["Invoice Date", "Invoice Number", "Reference", "VAT Number"]
    for i, l in enumerate(meta_lignes):
        if l in labels and i + 1 < len(meta_lignes):
            meta[l] = meta_lignes[i + 1]

    # Bloc fournisseur (droite)
    fournisseur = zone_lignes(page, w * 0.76, w, 0, bas)

    return client, meta, fournisseur


def extraire_lignes(texte):
    lignes, courant = [], None
    for l in texte.splitlines():
        l = l.strip()
        m = LIGNE_RE.match(l)
        if m:
            courant = {
                "Item": m["item"],
                "Description": m["desc"],
                "Quantité": to_float(m["qte"]),
                "Prix unitaire": to_float(m["pu"]),
                "TVA": m["tva"],
                "Montant": to_float(m["montant"]),
            }
            lignes.append(courant)
        elif courant is not None:
            if l.startswith("Subtotal"):
                break
            courant["Description"] += " " + l  # suite de la description
    return lignes


def extraire_facture(pdf_path):
    with pdfplumber.open(pdf_path) as pdf:
        page1 = pdf.pages[0]
        texte = "\n".join(p.extract_text(x_tolerance=1) or "" for p in pdf.pages)
        client, meta, fournisseur = extraire_entete(page1)

    devise = search(r"TOTAL\s+([A-Z]{3})", texte)
    entete = {
        "Fichier": Path(pdf_path).name,
        "Type": "TAX INVOICE" if "TAX INVOICE" in texte else None,
        "Numéro facture": meta.get("Invoice Number") or search(r"(INV-\d+)", texte),
        "Date facture": to_date(meta.get("Invoice Date")),
        "Date échéance": to_date(search(r"Due Date:\s*(.+)", texte)),
        "Référence": meta.get("Reference"),
        "N° TVA": meta.get("VAT Number"),
        "Devise": devise,
        "Client": client[0] if client else None,
        "Adresse client": ", ".join(l.rstrip(",") for l in client[1:]),
        "Fournisseur": fournisseur[0] if fournisseur else None,
        "Adresse fournisseur": ", ".join(
            l.rstrip(",") for l in fournisseur[1:] if not l.startswith(("tel", "email")) and "@" not in l
        ),
        "Téléphone fournisseur": search(r"tel\.:\s*(\S+)", texte),
        "Email fournisseur": search(r"([\w.+-]+@[\w-]+\.[\w.]+)", texte),
        "N° société (Company Reg.)": search(r"Company Registration No:\s*(\d+)", texte),
        "Sous-total": to_float(search(rf"Subtotal\s+({NUM})", texte)),
        "Total": to_float(search(rf"TOTAL\s+[A-Z]{{3}}\s+({NUM})", texte)),
        "Montant payé": to_float(search(rf"Less Amount Paid\s+({NUM})", texte)),
        "Montant dû": to_float(search(rf"AMOUNT DUE\s+[A-Z]{{3}}\s+({NUM})", texte)),
        "Banque": search(r"^(.*Bank Plc)$", texte, re.M),
        "N° compte": search(r"^Account Number:\s*(\d+)", texte, re.M),
        "Sort Code": search(r"Sort Code:\s*(\d+)", texte),
        "IBAN": search(r"IBAN:\s*(\S+)", texte),
        "BIC / SWIFT": search(r"BIC / SWIFT:\s*(\S+)", texte),
    }
    lignes = extraire_lignes(texte)

    # Contrôle de cohérence : somme des lignes = sous-total
    somme = round(sum(l["Montant"] for l in lignes), 2)
    entete["Contrôle (somme lignes = sous-total)"] = (
        "OK" if somme == entete["Sous-total"] else f"ÉCART : lignes = {somme}"
    )
    return entete, lignes


# ---------------------------------------------------------------- Excel
ENTETE_STYLE = dict(
    font=Font(bold=True, color="FFFFFF"),
    fill=PatternFill("solid", fgColor="1F4E78"),
    alignment=Alignment(horizontal="center", vertical="center"),
)


def styler_entete(ws, row=1):
    for cell in ws[row]:
        for k, v in ENTETE_STYLE.items():
            setattr(cell, k, v)


def ajuster_colonnes(ws, max_width=70):
    for col in ws.columns:
        largeur = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(largeur + 2, max_width)


def ecrire_excel(entete, lignes, sortie):
    wb = Workbook()

    # Feuille 1 : informations générales
    ws = wb.active
    ws.title = "Facture"
    ws.append(["Champ", "Valeur"])
    for k, v in entete.items():
        ws.append([k, v])
    styler_entete(ws)
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        c = row[0]
        if isinstance(c.value, float):
            c.number_format = "#,##0.00"
        elif hasattr(c.value, "year"):
            c.number_format = "DD/MM/YYYY"
    ajuster_colonnes(ws)

    # Feuille 2 : lignes de facture
    ws2 = wb.create_sheet("Lignes")
    colonnes = ["N° facture", "Item", "Description", "Quantité", "Prix unitaire", "TVA", "Montant"]
    ws2.append(colonnes)
    for l in lignes:
        ws2.append([entete["Numéro facture"]] + [l[c] for c in colonnes[1:]])
    n = len(lignes)
    ws2.append(["", "", "TOTAL", f"=SUM(D2:D{n + 1})", "", "", f"=SUM(G2:G{n + 1})"])
    for cell in ws2[n + 2]:
        cell.font = Font(bold=True)
    styler_entete(ws2)
    for row in ws2.iter_rows(min_row=2, min_col=4, max_col=7):
        for c in row:
            c.number_format = "#,##0.00"
    ajuster_colonnes(ws2)
    ws2.freeze_panes = "A2"

    wb.save(sortie)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("pdfs", nargs="*", default=[str(BASE_DIR / "Data" / "INV-55868.pdf")])
    parser.add_argument("-o", "--output", default=str(BASE_DIR / "Output"))
    args = parser.parse_args()

    out_dir = Path(args.output)
    out_dir.mkdir(parents=True, exist_ok=True)
    for pdf in args.pdfs:
        entete, lignes = extraire_facture(pdf)
        sortie = out_dir / f"{Path(pdf).stem}.xlsx"
        ecrire_excel(entete, lignes, sortie)
        print(f"{Path(pdf).name} -> {sortie}  ({len(lignes)} lignes, total {entete['Total']} {entete['Devise']}, "
              f"contrôle : {entete['Contrôle (somme lignes = sous-total)']})")


if __name__ == "__main__":
    main()
