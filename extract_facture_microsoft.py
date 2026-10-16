"""Extraction des informations d'une facture Microsoft (Azure / Billing Summary) vers Excel.

Usage :
    python extract_facture_microsoft.py [facture.pdf ...] [-o dossier_sortie]

Sans argument, traite Data/G133453773.pdf.
"""
import argparse
import re
from datetime import datetime
from pathlib import Path

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Font

from extract_facture import NUM, ajuster_colonnes, search, styler_entete, to_float

BASE_DIR = Path(__file__).resolve().parent
DATE = r"\d{2}/\d{2}/\d{4}"

# Achat : période, prix unitaire, quantité, charges, total HT, indicateur TVA
ACHAT_RE = re.compile(
    rf"^(?P<debut>{DATE})-(?P<fin>{DATE})\s+(?P<pu>{NUM})\s+(?P<qte>\d+)\s+"
    rf"(?P<charges>{NUM})\s+(?P<total>{NUM})\s+(?P<ind>[A-Z](?:,[A-Z])*)$"
)
# Consommation : libellé, charges, total HT, indicateur TVA (période sur la ligne suivante)
USAGE_RE = re.compile(
    rf"^(?P<libelle>.+?)\s+(?P<charges>{NUM})\s+(?P<total>{NUM})\s+(?P<ind>[A-Z](?:,[A-Z])*)$"
)
PERIODE_RE = re.compile(rf"^({DATE})\s*-\s*({DATE})$")
TAXE_RE = re.compile(rf"^(?P<desc>VAT\s+[\d.]+%)\s+(?P<montant>{NUM})\s+(?P<ind>[A-Z])$")
ENTETES_TABLEAU = ("Purchases", "Charge Start Date", "Date ")


def to_date(s):
    try:
        return datetime.strptime(s.strip(), "%d/%m/%Y").date()
    except (ValueError, AttributeError):
        return s


def bloc_adresse(page, x0, x1, top, bottom):
    zone = page.crop((x0, top, x1, bottom)).extract_text(x_tolerance=1) or ""
    return [l.strip() for l in zone.splitlines() if l.strip()]


def extraire_adresses(page):
    """Fournisseur (haut gauche), Sold To et Bill To, repérés par leur position."""
    w = page.width
    sold = page.search("Sold To")[0]
    bill = page.search("Bill To")[0]
    fin = page.search("This invoice is for")[0]["top"] - 2
    milieu = w * 0.49  # colonne de droite (Summary / Billing Profile)

    titre = page.search("Microsoft Ireland") or page.search("Microsoft")
    fournisseur = bloc_adresse(page, 0, milieu, titre[0]["top"] - 1, sold["top"] - 1)
    sold_to = bloc_adresse(page, 0, bill["x0"] - 2, sold["bottom"] + 1, fin)
    bill_to = bloc_adresse(page, bill["x0"] - 2, milieu, bill["bottom"] + 1, fin)
    return fournisseur, sold_to, bill_to


def extraire_sections(texte):
    """Noms des sections (profils de facturation) listés dans « Section Summary »."""
    m = re.search(r"Section Summary(.*?)Total \(excluding Tax\)", texte, re.S)
    if not m:
        return set()
    return {
        r.group(1).strip()
        for r in re.finditer(rf"^(.+?)\s+{NUM}\s+{NUM}$", m.group(1), re.M)
    }


def extraire_lignes(texte, sections):
    debut = texte.find("Billing Details By Product")
    if debut < 0:
        return []
    lignes, section, produit, precedente, derniere = [], None, None, None, None
    for l in texte[debut:].splitlines()[1:]:
        l = l.strip()
        if not l:
            continue
        if l.startswith(("Tax Description", "Tax Amount")):
            break
        if l in sections:
            section = l
        elif l.startswith(ENTETES_TABLEAU):
            if l.startswith("Purchases") and precedente:
                produit = precedente  # le titre du produit précède l'en-tête du tableau
        elif m := ACHAT_RE.match(l):
            derniere = {
                "Section": section, "Produit": produit, "Type": "Achat", "Libellé": produit,
                "Début": to_date(m["debut"]), "Fin": to_date(m["fin"]),
                "Prix unitaire": to_float(m["pu"]), "Qté": int(m["qte"]),
                "Charges": to_float(m["charges"]), "Total HT": to_float(m["total"]),
                "Indicateur TVA": m["ind"],
            }
            lignes.append(derniere)
        elif m := USAGE_RE.match(l):
            derniere = {
                "Section": section, "Produit": produit, "Type": "Consommation",
                "Libellé": m["libelle"], "Début": None, "Fin": None,
                "Prix unitaire": None, "Qté": None,
                "Charges": to_float(m["charges"]), "Total HT": to_float(m["total"]),
                "Indicateur TVA": m["ind"],
            }
            lignes.append(derniere)
        elif (m := PERIODE_RE.match(l)) and derniere and derniere["Début"] is None:
            derniere["Début"], derniere["Fin"] = to_date(m[1]), to_date(m[2])
        precedente = l
    return lignes


def extraire_facture(pdf_path):
    with pdfplumber.open(pdf_path) as pdf:
        pages = [p.extract_text(x_tolerance=1) or "" for p in pdf.pages]
        # La page « Invoice » contient les adresses ; à défaut, la première page
        idx = next((i for i, t in enumerate(pages) if t.lstrip().startswith("Invoice")), 0)
        fournisseur, sold_to, bill_to = extraire_adresses(pdf.pages[idx])
    texte = "\n".join(pages)

    tva_fournisseur = next((l for l in fournisseur if l.startswith("VAT Reg")), None)
    fournisseur = [l for l in fournisseur if l != tva_fournisseur]
    periode = re.search(rf"billing period\s+({DATE})\s*-\s*({DATE})", texte)
    total = re.search(rf"^Total\s+([A-Z]{{3}})\s+({NUM})$", texte, re.M)

    entete = {
        "Fichier": Path(pdf_path).name,
        "Fournisseur": fournisseur[0] if fournisseur else None,
        "Adresse fournisseur": ", ".join(fournisseur[1:]),
        "N° TVA fournisseur": search(r"VAT Reg\. No\.\s*(.+)", texte),
        "Numéro facture": search(r"Tax Invoice Number\s+(\S+)", texte) or search(r"Billing Number\s+(\S+)", texte),
        "Date facture": to_date(search(rf"Tax Invoice Date\s+({DATE})", texte)
                                or search(rf"Document Date\s+({DATE})", texte)),
        "Date échéance": to_date(search(rf"Due on\s+({DATE})", texte)),
        "Période début": to_date(periode[1]) if periode else None,
        "Période fin": to_date(periode[2]) if periode else None,
        "Profil de facturation": search(r"Billing Profile\s+(.+)", texte),
        "Vendu à (Sold To)": sold_to[0] if sold_to else None,
        "Adresse Sold To": ", ".join(sold_to[1:]),
        "Facturé à (Bill To)": bill_to[0] if bill_to else None,
        "Adresse Bill To": ", ".join(bill_to[1:]),
        "Devise": total[1] if total else None,
        "Charges": to_float(search(rf"^Charges\s+({NUM})$", texte, re.M)),
        "Sous-total": to_float(search(rf"Subtotal\s+({NUM})", texte)),
        "Crédit Azure": to_float(search(r"Azure Credit\s+([\d,.]+)", texte)),
        "Taxe": to_float(search(rf"^Tax\s+({NUM})$", texte, re.M)),
        "Total": to_float(total[2]) if total else None,
        "Instructions de paiement": search(r"Payment Instructions:\s*\n(.+)", texte),
    }

    lignes = extraire_lignes(texte, extraire_sections(texte))
    taxes = [
        {"Description": m["desc"], "Montant": to_float(m["montant"]), "Indicateur": m["ind"]}
        for m in (TAXE_RE.match(l.strip()) for l in texte.splitlines()) if m
    ]

    somme = round(sum(l["Total HT"] for l in lignes), 2)
    entete["Contrôle (somme lignes = sous-total)"] = (
        "OK" if somme == entete["Sous-total"] else f"ÉCART : lignes = {somme}"
    )
    return entete, lignes, taxes


def ecrire_tableau(ws, colonnes, lignes, formats):
    ws.append(colonnes)
    for l in lignes:
        ws.append([l[c] for c in colonnes])
    styler_entete(ws)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            fmt = formats.get(colonnes[c.column - 1])
            if fmt:
                c.number_format = fmt
    ws.freeze_panes = "A2"


def ecrire_excel(entete, lignes, taxes, sortie):
    wb = Workbook()

    ws = wb.active
    ws.title = "Facture"
    ws.append(["Champ", "Valeur"])
    for k, v in entete.items():
        ws.append([k, v])
    styler_entete(ws)
    for (c,) in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        if isinstance(c.value, float):
            c.number_format = "#,##0.00"
        elif hasattr(c.value, "year"):
            c.number_format = "DD/MM/YYYY"
    ajuster_colonnes(ws)

    montants = {k: "#,##0.00" for k in ("Prix unitaire", "Charges", "Total HT", "Montant")}
    dates = {"Début": "DD/MM/YYYY", "Fin": "DD/MM/YYYY"}

    ws2 = wb.create_sheet("Lignes")
    colonnes = ["Section", "Produit", "Type", "Libellé", "Début", "Fin",
                "Prix unitaire", "Qté", "Charges", "Total HT", "Indicateur TVA"]
    ecrire_tableau(ws2, colonnes, lignes, {**montants, **dates})
    n = len(lignes) + 1
    ws2.append(["", "", "", "TOTAL", "", "", "", "", f"=SUM(I2:I{n})", f"=SUM(J2:J{n})", ""])
    for cell in ws2[n + 1]:
        cell.font = Font(bold=True)
        if cell.column in (9, 10):
            cell.number_format = "#,##0.00"
    ajuster_colonnes(ws2)

    ws3 = wb.create_sheet("Taxes")
    ecrire_tableau(ws3, ["Description", "Montant", "Indicateur"], taxes, montants)
    ajuster_colonnes(ws3)

    wb.save(sortie)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("pdfs", nargs="*", default=[str(BASE_DIR / "Data" / "G133453773.pdf")])
    parser.add_argument("-o", "--output", default=str(BASE_DIR / "Output"))
    args = parser.parse_args()

    out_dir = Path(args.output)
    out_dir.mkdir(parents=True, exist_ok=True)
    for pdf in args.pdfs:
        entete, lignes, taxes = extraire_facture(pdf)
        sortie = out_dir / f"{Path(pdf).stem}.xlsx"
        ecrire_excel(entete, lignes, taxes, sortie)
        print(f"{Path(pdf).name} -> {sortie}  ({len(lignes)} lignes, total {entete['Total']} {entete['Devise']}, "
              f"contrôle : {entete['Contrôle (somme lignes = sous-total)']})")


if __name__ == "__main__":
    main()
