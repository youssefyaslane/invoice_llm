"""Extrait toutes les factures PDF d'un dossier (sous-dossiers compris) avec le LLM local
et regroupe les résultats dans un fichier Excel, avec le chemin du PDF source.

Usage (voir docker-compose.yml) :
    docker compose run --rm main                          -> Data/2026 -> Output/factures_2026.xlsx
    docker compose run --rm main Data/2026/01-JAN         -> un seul mois
    docker compose run --rm main --refaire                -> refait aussi les factures déjà extraites
    docker compose run --rm main --refaire ReadAI 0008    -> refait seulement celles dont le chemin contient "ReadAI" ou "0008"
ou en local :
    .venv-llm/bin/python main.py [dossier] [-m modele] [--url http://localhost:11434]

Chaque résultat est gardé dans Output/resultats/ : une facture déjà extraite n'est pas renvoyée au LLM
(sauf --refaire ou PDF modifié), donc relancer le script chaque mois ne traite que les nouvelles factures.
"""
import argparse
import json
import os
import time
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from test_llm_facture import BASE_DIR, extraire_facture, verifier

DATA = BASE_DIR / "Data"
A_VERIFIER = PatternFill("solid", fgColor="FFC7CE")
ENTETE = PatternFill("solid", fgColor="DDEBF7")


def source(pdf):
    """Chemin du PDF à partir de Data/, comme dans OneDrive : 2026/01-JAN/OVH/1 Jan 2026/Facture_FR74551374.pdf"""
    try:
        return pdf.resolve().relative_to(DATA.resolve()).as_posix()
    except ValueError:
        return pdf.as_posix()


def traiter(pdf, args, cache):
    """Résultat d'une facture : repris de Output/resultats/ s'il est à jour, sinon extrait par le LLM."""
    refaire = args.refaire is not None and (not args.refaire or any(t.lower() in source(pdf).lower() for t in args.refaire))
    if cache.is_file() and cache.stat().st_mtime > pdf.stat().st_mtime and not refaire:
        print("  déjà extraite, résultat repris")
        return json.loads(cache.read_text(encoding="utf-8"))

    texte, data, rep, duree = extraire_facture(pdf, args.url, args.modele)
    resultat = {"data": data, "controles": verifier(data, texte), "duree": round(duree, 1)}
    cache.parent.mkdir(parents=True, exist_ok=True)
    cache.write_text(json.dumps(resultat, indent=2, ensure_ascii=False), encoding="utf-8")
    return resultat


def date_excel(texte):
    try:
        return datetime.strptime(texte, "%d/%m/%Y")
    except (TypeError, ValueError):
        return texte


def feuille(wb, titre, colonnes, lignes, formats, rouges=()):
    ws = wb.create_sheet(titre)
    ws.append(colonnes)
    for ligne in lignes:
        ws.append([ligne.get(c) for c in colonnes])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True), ENTETE
    for i, row in enumerate(ws.iter_rows(min_row=2), start=2):
        for c in row:
            fmt = formats.get(colonnes[c.column - 1])
            if fmt:
                c.number_format = fmt
            if i - 2 in rouges:
                c.fill = A_VERIFIER
    for i, nom in enumerate(colonnes, start=1):
        largeur = max([len(str(nom))] + [len(str(ws.cell(r, i).value or "")) for r in range(2, ws.max_row + 1)])
        ws.column_dimensions[ws.cell(1, i).column_letter].width = min(largeur + 2, 70)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def ecrire_excel(resultats, sortie):
    factures, articles, rouges = [], [], set()
    for i, (pdf, r) in enumerate(resultats):
        data, ctrl = r.get("data") or {}, r.get("controles") or {}
        problemes = [f"{nom} : {v}" for nom, v in ctrl.items() if v != "OK"]
        if r.get("erreur"):
            problemes = [f"erreur : {r['erreur']}"]
        if problemes:
            rouges.add(i)
        commun = {"Source PDF": source(pdf), "Fournisseur": data.get("fournisseur"),
                  "Date paiement": date_excel(data.get("date_paiement")), "Devise": data.get("devise")}
        factures.append({**commun, "Total": data.get("total"), "Nb articles": len(data.get("articles") or []),
                         "À vérifier": "Oui" if problemes else "Non", "Contrôles": " | ".join(problemes) or "OK",
                         "Durée LLM (s)": r.get("duree")})
        for a in data.get("articles") or []:
            articles.append({**commun, "Section": a.get("section"), "Description": a.get("description"),
                             "Quantité": a.get("quantite"),
                             "Prix unitaire": a.get("prix_unitaire"), "Prix": a.get("prix")})

    wb = Workbook()
    wb.remove(wb.active)
    montant, date = "#,##0.00", "DD/MM/YYYY"
    feuille(wb, "Factures", ["Source PDF", "Fournisseur", "Date paiement", "Devise", "Total", "Nb articles",
                             "À vérifier", "Contrôles", "Durée LLM (s)"],
            factures, {"Date paiement": date, "Total": montant}, rouges)
    feuille(wb, "Articles", ["Source PDF", "Fournisseur", "Date paiement", "Devise", "Section", "Description", "Quantité",
                             "Prix unitaire", "Prix"],
            articles, {"Date paiement": date, "Prix unitaire": montant, "Prix": montant})
    sortie.parent.mkdir(parents=True, exist_ok=True)
    wb.save(sortie)
    return len(rouges)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("dossier", nargs="?", default=DATA / "2026", type=Path)
    ap.add_argument("-m", "--modele", default=os.environ.get("MODELE", "qwen2.5:3b"))
    ap.add_argument("--url", default=os.environ.get("OLLAMA_URL", "http://localhost:11434"))
    ap.add_argument("--sortie", type=Path, help="fichier Excel (par défaut Output/factures_<dossier>.xlsx)")
    ap.add_argument("--refaire", nargs="*", metavar="TEXTE",
                    help="renvoie au LLM les factures déjà extraites (toutes, ou celles dont le chemin contient un des TEXTE)")
    args = ap.parse_args()

    pdfs = sorted(p for p in args.dossier.rglob("*") if p.suffix.lower() == ".pdf")
    if not pdfs:
        raise SystemExit(f"Aucun PDF dans {args.dossier}")
    sortie = args.sortie or BASE_DIR / "Output" / f"factures_{args.dossier.name}.xlsx"
    print(f"{len(pdfs)} factures dans {args.dossier}  —  modèle : {args.modele}\n")

    t0, resultats = time.time(), []
    for i, pdf in enumerate(pdfs, start=1):
        print(f"[{i}/{len(pdfs)}] {source(pdf)}")
        cache = BASE_DIR / "Output" / "resultats" / Path(source(pdf)).with_suffix(".json")
        try:
            r = traiter(pdf, args, cache)
        except Exception as e:  # une facture illisible ne doit pas arrêter les autres
            print(f"  ERREUR : {e}")
            r = {"erreur": str(e)}
        data = r.get("data") or {}
        print(f"  {data.get('fournisseur')}  {data.get('date_paiement')}  {data.get('total')} {data.get('devise')}"
              f"  ({len(data.get('articles') or [])} articles)\n")
        resultats.append((pdf, r))

    a_verifier = ecrire_excel(resultats, sortie)
    print(f"Terminé en {(time.time() - t0) / 60:.1f} min : {len(pdfs)} factures, {a_verifier} à vérifier")
    print(f"Excel : {sortie}")


if __name__ == "__main__":
    main()
